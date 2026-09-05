from __future__ import annotations

import math
import re
import textwrap
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "entrega-2"
ASSET_DIR = OUT_DIR / "assets"
DOCX_PATH = OUT_DIR / "academia-conectada-entrega-2.docx"

BRAND = "#b90812"
INK = "#20242c"
MUTED = "#64707d"
PAPER = "#f7f7f2"
LINE = "#d8ded6"
TEAL = "#127c86"
GREEN = "#23824a"
AMBER = "#f3b43f"
BLUE = "#2f5f98"


def ensure_dirs() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
        if bold
        else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Helvetica Bold.ttf"
        if bold
        else "/System/Library/Fonts/Supplemental/Helvetica.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def text_width(draw: ImageDraw.ImageDraw, value: str, fnt: ImageFont.ImageFont) -> int:
    if not value:
        return 0
    bbox = draw.textbbox((0, 0), value, font=fnt)
    return bbox[2] - bbox[0]


def wrap_lines(
    draw: ImageDraw.ImageDraw,
    value: str,
    fnt: ImageFont.ImageFont,
    max_width: int,
) -> list[str]:
    words = value.split()
    lines: list[str] = []
    current = ""
    for word in words:
        probe = word if not current else f"{current} {word}"
        if text_width(draw, probe, fnt) <= max_width:
            current = probe
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]


def draw_wrapped(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    value: str,
    fnt: ImageFont.ImageFont,
    fill: str,
    max_width: int,
    line_gap: int = 8,
) -> int:
    x, y = xy
    line_height = fnt.size + line_gap if hasattr(fnt, "size") else 22
    for line in wrap_lines(draw, value, fnt, max_width):
        draw.text((x, y), line, font=fnt, fill=fill)
        y += line_height
    return y


def centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    value: str,
    fnt: ImageFont.ImageFont,
    fill: str = INK,
) -> None:
    x1, y1, x2, y2 = box
    lines = wrap_lines(draw, value, fnt, x2 - x1 - 28)
    line_height = fnt.size + 5 if hasattr(fnt, "size") else 22
    y = y1 + ((y2 - y1) - len(lines) * line_height) // 2
    for line in lines:
        w = text_width(draw, line, fnt)
        draw.text((x1 + (x2 - x1 - w) // 2, y), line, font=fnt, fill=fill)
        y += line_height


def rounded_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    fill: str,
    outline: str = LINE,
    radius: int = 12,
    width: int = 2,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    fill: str = "#485260",
    width: int = 3,
) -> None:
    draw.line([start, end], fill=fill, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 12
    points = [
        end,
        (
            int(end[0] - size * math.cos(angle - math.pi / 6)),
            int(end[1] - size * math.sin(angle - math.pi / 6)),
        ),
        (
            int(end[0] - size * math.cos(angle + math.pi / 6)),
            int(end[1] - size * math.sin(angle + math.pi / 6)),
        ),
    ]
    draw.polygon(points, fill=fill)


def draw_button(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    label: str,
    fill: str = BRAND,
    text_fill: str = "white",
) -> None:
    rounded_box(draw, box, fill=fill, outline=fill, radius=9)
    centered_text(draw, box, label, font(23, True), text_fill)


def save_image(image: Image.Image, name: str) -> Path:
    path = ASSET_DIR / name
    image.save(path)
    return path


def draw_app_shell(title: str, subtitle: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    img = Image.new("RGB", (1600, 1000), PAPER)
    draw = ImageDraw.Draw(img)
    rounded_box(draw, (48, 46, 1552, 112), "white", "#e2e6df", radius=14)
    rounded_box(draw, (74, 62, 124, 96), BRAND, BRAND, radius=8)
    centered_text(draw, (74, 62, 124, 96), "AC", font(20, True), "white")
    draw.text((146, 58), "Academia Conectada", font=font(24, True), fill=BRAND)
    draw.text((146, 86), subtitle, font=font(17), fill=MUTED)
    draw.text((1210, 69), title, font=font(24, True), fill=INK)
    rounded_box(draw, (48, 142, 270, 930), "white", "#e2e6df", radius=14)
    nav = ["Visão", "Planos", "Aulas", "Treino", "Gestão", "Admin"]
    for index, item in enumerate(nav):
        y = 180 + index * 72
        fill = "#fff4f4" if index < 4 else "#f6f7f3"
        outline = BRAND if index == 0 else "#e4e8e0"
        rounded_box(draw, (76, y, 242, y + 48), fill, outline, radius=10)
        draw.text((104, y + 13), item, font=font(20, True), fill=BRAND if index == 0 else "#4a5562")
    return img, draw


def prototype_public() -> Path:
    img, draw = draw_app_shell("Perfil Visitante", "Planos, modalidades e horários")
    draw.text((310, 164), "Escolha de plano com matrícula demonstrativa", font=font(34, True), fill=INK)
    draw_wrapped(
        draw,
        (310, 210),
        "Tela pública para consultar planos, modalidades, disponibilidade e iniciar cadastro.",
        font(22),
        MUTED,
        840,
    )
    plan_data = [
        ("Essencial", "R$ 99", "Musculação e treino inicial"),
        ("Conectado", "R$ 149", "Aulas coletivas e reserva online"),
        ("Performance", "R$ 219", "Acompanhamento personalizado"),
    ]
    for index, (name, price, text) in enumerate(plan_data):
        x = 310 + index * 390
        rounded_box(draw, (x, 292, x + 345, 580), "white", BRAND if index == 1 else LINE, radius=14, width=3)
        draw.text((x + 28, 324), name, font=font(30, True), fill=INK)
        draw.text((x + 28, 370), price, font=font(38, True), fill=BRAND)
        draw_wrapped(draw, (x + 28, 430), text, font(21), MUTED, 280)
        draw_button(draw, (x + 28, 500, x + 286, 548), "Solicitar matrícula", fill=BRAND if index == 1 else INK)
    draw.text((310, 650), "Aulas em destaque", font=font(30, True), fill=INK)
    for index, item in enumerate(["Bike indoor - 19:00", "Funcional - 07:30", "Pilates solo - lotada"]):
        y = 704 + index * 64
        rounded_box(draw, (310, y, 1380, y + 48), "white", LINE, radius=10)
        draw.text((338, y + 13), item, font=font(21, True), fill=INK)
        draw.text((1180, y + 13), "Ver detalhes", font=font(20, True), fill=TEAL)
    return save_image(img, "figura-02-prototipo-visitante.png")


def prototype_student() -> Path:
    img, draw = draw_app_shell("Perfil Aluno", "Matrícula, aulas e treino")
    draw.text((310, 164), "Dashboard do aluno", font=font(36, True), fill=INK)
    stats = [("Plano", "Conectado"), ("Reservas", "1"), ("Treino", "50%")]
    for index, (label, value) in enumerate(stats):
        x = 310 + index * 290
        rounded_box(draw, (x, 238, x + 248, 356), "white", LINE, radius=14)
        draw.text((x + 24, 262), label, font=font(20, True), fill=MUTED)
        draw.text((x + 24, 296), value, font=font(34, True), fill=INK)
    rounded_box(draw, (310, 410, 920, 850), "white", LINE, radius=14)
    draw.text((342, 444), "Sessão A - registro de treino", font=font(29, True), fill=INK)
    exercises = ["Leg press - 4 x 10 - 80 kg", "Supino reto - 4 x 8 - 42 kg", "Remada baixa - 3 x 12", "Prancha - 3 x 40 s"]
    for index, item in enumerate(exercises):
        y = 508 + index * 70
        rounded_box(draw, (342, y, 884, y + 48), "#fdfdfb", LINE, radius=9)
        draw.ellipse((362, y + 13, 384, y + 35), outline=GREEN, width=3, fill="#e3f4e8" if index < 2 else "white")
        if index < 2:
            draw.line((366, y + 24, 374, y + 32, 382, y + 18), fill=GREEN, width=3)
        draw.text((402, y + 13), item, font=font(20, True), fill=INK)
    rounded_box(draw, (980, 410, 1410, 850), "#f6f9fc", "#dbe5ef", radius=14)
    draw.text((1014, 444), "Evolução mensal", font=font(28, True), fill=INK)
    for index, value in enumerate([44, 52, 59, 67, 76]):
        x = 1034 + index * 70
        draw.rounded_rectangle((x, 760 - value * 3, x + 46, 760), radius=7, fill=TEAL)
        draw.text((x + 4, 778), ["Mai", "Jun", "Jul", "Ago", "Set"][index], font=font(16, True), fill=MUTED)
    draw_button(draw, (1014, 804, 1368, 858), "Salvar registro", fill=INK)
    return save_image(img, "figura-03-prototipo-aluno.png")


def prototype_management() -> Path:
    img, draw = draw_app_shell("Instrutor e Admin", "Gestão operacional")
    draw.text((310, 164), "Gestão de treinos e cadastros", font=font(36, True), fill=INK)
    rounded_box(draw, (310, 238, 920, 860), "white", LINE, radius=14)
    draw.text((342, 274), "Área do instrutor", font=font(29, True), fill=INK)
    rows = [("Lara Souza", "Hipertrofia A/B", "Ativo"), ("Caio Martins", "Condicionamento", "Revisão"), ("Nina Alves", "Retorno gradual", "Aguardando")]
    for index, (name, plan, status) in enumerate(rows):
        y = 344 + index * 104
        rounded_box(draw, (342, y, 884, y + 72), "#fdfdfb", LINE, radius=10)
        draw.text((368, y + 14), name, font=font(23, True), fill=INK)
        draw.text((368, y + 44), plan, font=font(18), fill=MUTED)
        draw.text((724, y + 25), status, font=font(18, True), fill=BRAND if status == "Revisão" else TEAL)
    rounded_box(draw, (980, 238, 1410, 860), "white", LINE, radius=14)
    draw.text((1014, 274), "Administração", font=font(29, True), fill=INK)
    metrics = [("Alunos", "128"), ("Planos", "3"), ("Aulas", "24"), ("Ocupação", "78%")]
    for index, (label, value) in enumerate(metrics):
        x = 1014 + (index % 2) * 182
        y = 344 + (index // 2) * 120
        rounded_box(draw, (x, y, x + 156, y + 86), "#f7f7f2", LINE, radius=10)
        draw.text((x + 18, y + 16), label, font=font(18, True), fill=MUTED)
        draw.text((x + 18, y + 42), value, font=font(29, True), fill=INK)
    for index, row in enumerate(["Planos", "Modalidades", "Aulas coletivas", "Instrutores"]):
        y = 620 + index * 48
        draw.text((1014, y), row, font=font(20, True), fill=INK)
        draw.text((1280, y), "Manter", font=font(19, True), fill=BRAND)
    return save_image(img, "figura-04-prototipo-gestao.png")


def use_case_diagram() -> Path:
    img = Image.new("RGB", (1600, 1050), "white")
    draw = ImageDraw.Draw(img)
    draw.text((480, 44), "Diagrama geral de casos de uso", font=font(42, True), fill=INK)
    rounded_box(draw, (360, 140, 1240, 900), "#fbfcfb", "#4b5563", radius=4, width=3)
    draw.text((390, 172), "Sistema Academia Conectada", font=font(24, True), fill=INK)
    use_cases = [
        ("UC01\nCadastrar-se e entrar", 450, 260),
        ("UC02\nConsultar planos e modalidades", 820, 260),
        ("UC03\nSolicitar matrícula", 450, 390),
        ("UC04\nAgendar ou cancelar aula", 820, 390),
        ("UC05\nConsultar plano de treino", 450, 520),
        ("UC06\nRegistrar treino e evolução", 820, 520),
        ("UC07\nGerenciar treinos de alunos", 450, 650),
        ("UC08\nAdministrar cadastros, planos e aulas", 820, 650),
    ]
    centers: dict[str, tuple[int, int]] = {}
    for label, x, y in use_cases:
        draw.ellipse((x, y, x + 300, y + 70), outline="#475569", width=3, fill="white")
        centered_text(draw, (x + 10, y + 4, x + 290, y + 66), label, font(18, True), INK)
        centers[label.split("\n")[0]] = (x + 150, y + 35)

    def actor(x: int, y: int, label: str, color: str) -> tuple[int, int]:
        draw.ellipse((x - 18, y - 54, x + 18, y - 18), outline=color, width=4)
        draw.line((x, y - 18, x, y + 50), fill=color, width=4)
        draw.line((x - 44, y + 4, x + 44, y + 4), fill=color, width=4)
        draw.line((x, y + 50, x - 36, y + 100), fill=color, width=4)
        draw.line((x, y + 50, x + 36, y + 100), fill=color, width=4)
        w = text_width(draw, label, font(20, True))
        draw.text((x - w // 2, y + 118), label, font=font(20, True), fill=color)
        return (x, y + 22)

    visitante = actor(205, 315, "Visitante", BLUE)
    aluno = actor(205, 610, "Aluno", GREEN)
    instrutor = actor(1395, 520, "Instrutor", "#a64b19")
    admin = actor(1395, 740, "Administrador", BRAND)

    links = [
        (visitante, centers["UC01"], BLUE),
        (visitante, centers["UC02"], BLUE),
        (aluno, centers["UC01"], GREEN),
        (aluno, centers["UC02"], GREEN),
        (aluno, centers["UC03"], GREEN),
        (aluno, centers["UC04"], GREEN),
        (aluno, centers["UC05"], GREEN),
        (aluno, centers["UC06"], GREEN),
        (instrutor, centers["UC05"], "#a64b19"),
        (instrutor, centers["UC07"], "#a64b19"),
        (admin, centers["UC01"], BRAND),
        (admin, centers["UC02"], BRAND),
        (admin, centers["UC08"], BRAND),
    ]
    for start, end, color in links:
        draw.line([start, end], fill=color, width=2)
    draw.text((390, 858), "Associação: o ator participa ou inicia o caso de uso.", font=font(18), fill=MUTED)
    return save_image(img, "figura-01-casos-de-uso.png")


def domain_model() -> Path:
    img = Image.new("RGB", (1600, 1000), "white")
    draw = ImageDraw.Draw(img)
    draw.text((520, 48), "Modelo de domínio", font=font(42, True), fill=INK)

    groups = [
        ("Acesso e perfis", (70, 135, 1530, 320), "#f8fbff"),
        ("Contratação e aulas", (70, 360, 760, 740), "#fffaf0"),
        ("Treinos e evolução", (840, 360, 1530, 740), "#f6fbf7"),
        ("Relacionamentos principais", (70, 780, 1530, 940), "#fbfcfb"),
    ]
    for title, box, fill in groups:
        rounded_box(draw, box, fill, "#d5dde4", radius=16, width=2)
        draw.text((box[0] + 24, box[1] + 20), title, font=font(25, True), fill=INK)

    def entity(label: str, box: tuple[int, int, int, int], color: str = "#475569") -> tuple[int, int]:
        rounded_box(draw, box, "white", color, radius=10, width=3)
        centered_text(draw, box, label, font(21, True), INK)
        return ((box[0] + box[2]) // 2, (box[1] + box[3]) // 2)

    usuario = entity("Usuário", (660, 198, 940, 260), BRAND)
    aluno = entity("Aluno", (210, 245, 450, 300), GREEN)
    instrutor = entity("Instrutor", (680, 245, 920, 300), TEAL)
    admin = entity("Administrador", (1130, 245, 1410, 300), BLUE)
    draw.line((usuario[0], 260, usuario[0], 292), fill="#66717d", width=3)
    draw.line((330, 292, 1270, 292), fill="#66717d", width=3)
    for point in [aluno, instrutor, admin]:
        arrow(draw, (point[0], 292), (point[0], 302), "#66717d", 3)

    plano = entity("Plano", (120, 445, 310, 505), BRAND)
    matricula = entity("Matrícula", (390, 445, 610, 505), BRAND)
    aluno_ref = entity("Aluno", (250, 640, 480, 700), GREEN)
    aula = entity("Aula", (120, 555, 310, 615), TEAL)
    reserva = entity("Reserva", (390, 555, 610, 615), TEAL)
    arrow(draw, (320, 485), (385, 485), BRAND, 3)
    arrow(draw, (500, 505), (430, 640), BRAND, 3)
    arrow(draw, (320, 570), (385, 570), TEAL, 3)
    arrow(draw, (500, 615), (430, 640), TEAL, 3)

    instrutor_ref = entity("Instrutor", (900, 455, 1110, 515), TEAL)
    plano_treino = entity("Plano de treino", (1180, 455, 1460, 515), GREEN)
    sessao = entity("Sessão de treino", (900, 575, 1160, 635), GREEN)
    exercicio = entity("Exercício", (1225, 575, 1460, 635), GREEN)
    registro = entity("Registro de evolução", (1030, 660, 1340, 720), BLUE)
    arrow(draw, (1110, 485), (1180, 485), TEAL, 3)
    arrow(draw, (1320, 515), (1030, 575), GREEN, 3)
    arrow(draw, (1160, 605), (1225, 605), GREEN, 3)
    arrow(draw, (1030, 690), (955, 635), BLUE, 3)

    relations = [
        "1. Usuário especializa Aluno, Instrutor ou Administrador.",
        "2. Aluno possui Matrícula em um Plano e pode fazer Reservas em Aulas.",
        "3. Instrutor prescreve Plano de treino composto por Sessões e Exercícios.",
        "4. Aluno registra evolução vinculada ao treino e ao histórico de execução.",
    ]
    x = 110
    y = 838
    for index, relation in enumerate(relations):
        current_x = x if index < 2 else 830
        current_y = y + (index % 2) * 48
        draw_wrapped(draw, (current_x, current_y), relation, font(21, True), INK, 640, 6)
    return save_image(img, "figura-05-modelo-dominio.png")


def uml_class(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], title: str, attrs: Iterable[str], methods: Iterable[str]) -> None:
    x1, y1, x2, y2 = box
    rounded_box(draw, box, "white", "#475569", radius=8, width=3)
    draw.rectangle((x1, y1, x2, y1 + 42), fill="#edf2f7", outline="#475569", width=2)
    centered_text(draw, (x1, y1, x2, y1 + 42), title, font(18, True), INK)
    draw.line((x1, y1 + 118, x2, y1 + 118), fill="#475569", width=2)
    y = y1 + 54
    for attr in attrs:
        draw.text((x1 + 14, y), attr, font=font(15), fill=INK)
        y += 21
    y = y1 + 130
    for method in methods:
        draw.text((x1 + 14, y), method, font=font(15), fill=INK)
        y += 21


def class_diagram() -> Path:
    img = Image.new("RGB", (1700, 1150), "white")
    draw = ImageDraw.Draw(img)
    draw.text((555, 42), "Diagrama de classes de projeto", font=font(42, True), fill=INK)
    classes = [
        ((70, 140, 380, 320), "Usuario", ["+ id: string", "+ nome: string", "+ email: string"], ["+ autenticar()", "+ alterarPerfil()"]),
        ((70, 385, 380, 565), "Aluno", ["+ matriculaAtiva: boolean", "+ dataNascimento: Date"], ["+ solicitarMatricula()", "+ registrarTreino()"]),
        ((70, 615, 380, 795), "Instrutor", ["+ cref: string", "+ especialidade: string"], ["+ criarPlanoTreino()", "+ revisarPlano()"]),
        ((70, 845, 380, 1025), "Administrador", ["+ nivelAcesso: string"], ["+ manterCadastro()", "+ publicarAgenda()"]),
        ((510, 140, 840, 330), "Plano", ["+ nome: string", "+ valorMensal: decimal", "+ ativo: boolean"], ["+ ativar()", "+ desativar()"]),
        ((510, 385, 840, 575), "Matricula", ["+ protocolo: string", "+ status: string", "+ criadaEm: Date"], ["+ confirmar()", "+ encerrar()"]),
        ((510, 630, 840, 820), "Aula", ["+ dataHora: Date", "+ capacidade: number", "+ modalidade: string"], ["+ reservarVaga()", "+ cancelarReserva()"]),
        ((510, 875, 840, 1065), "Reserva", ["+ status: string", "+ criadaEm: Date"], ["+ confirmar()", "+ cancelar()"]),
        ((970, 140, 1315, 350), "PlanoTreino", ["+ objetivo: string", "+ vigenciaInicio: Date", "+ vigenciaFim: Date"], ["+ publicar()", "+ criarNovaVersao()"]),
        ((970, 425, 1315, 615), "Exercicio", ["+ nome: string", "+ grupoMuscular: string"], ["+ atualizarOrientacao()"]),
        ((970, 690, 1315, 900), "RegistroEvolucao", ["+ peso: decimal", "+ observacoes: string"], ["+ calcularHistorico()", "+ anexarMedida()"]),
    ]
    for args in classes:
        uml_class(draw, *args)

    arrow(draw, (225, 320), (225, 385), "#67707d", 3)
    arrow(draw, (225, 320), (225, 615), "#67707d", 3)
    arrow(draw, (225, 320), (225, 845), "#67707d", 3)
    draw.text((250, 350), "herança", font=font(16, True), fill=MUTED)
    draw.text((250, 580), "herança", font=font(16, True), fill=MUTED)
    draw.text((250, 810), "herança", font=font(16, True), fill=MUTED)

    rounded_box(draw, (1375, 140, 1635, 900), "#fbfcfb", "#d5dde4", radius=16, width=2)
    draw.text((1405, 178), "Associações", font=font(24, True), fill=INK)
    association_lines = [
        "Aluno 1..* Matrícula",
        "Matrícula 1 Plano",
        "Aluno 0..* Reserva",
        "Reserva 1 Aula",
        "Instrutor 0..* PlanoTreino",
        "PlanoTreino 1..* Exercício",
        "Aluno 0..* RegistroEvolucao",
    ]
    y = 230
    for line in association_lines:
        draw_wrapped(draw, (1405, y), line, font(19, True), "#3d4652", 200, 6)
        y += 72

    rounded_box(draw, (955, 950, 1635, 1075), "#f6f9fc", "#d5dde4", radius=16, width=2)
    draw.text((990, 982), "Serviços previstos", font=font(22, True), fill=INK)
    draw_wrapped(
        draw,
        (990, 1020),
        "AuthService, MatriculaService, AgendaService, TreinoService e AdminService coordenam regras de negócio e repositórios.",
        font(18, True),
        MUTED,
        590,
        5,
    )
    return save_image(img, "figura-06-diagrama-classes.png")


def sequence_diagram(name: str, title: str, actors: list[str], steps: list[tuple[int, int, str]]) -> Path:
    width = 1550
    height = 880
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 42), title, font=font(38, True), fill=INK)
    xs = [120 + index * ((width - 240) // (len(actors) - 1)) for index in range(len(actors))]
    for x, actor_name in zip(xs, actors):
        rounded_box(draw, (x - 105, 116, x + 105, 166), "#edf2f7", "#475569", radius=8, width=2)
        centered_text(draw, (x - 105, 116, x + 105, 166), actor_name, font(18, True), INK)
        draw.line((x, 166, x, height - 70), fill="#9aa3ad", width=2)
    y = 230
    for from_idx, to_idx, label in steps:
        start = (xs[from_idx], y)
        end = (xs[to_idx], y)
        arrow(draw, start, end, BRAND if from_idx < to_idx else TEAL, 3)
        lx = min(start[0], end[0]) + 20
        draw.rectangle((lx, y - 32, lx + 455, y - 6), fill="white")
        draw.text((lx + 6, y - 31), label, font=font(17, True), fill=INK)
        y += 82
    return save_image(img, name)


def sequence_images() -> list[Path]:
    enrollment = sequence_diagram(
        "figura-07-sequencia-matricula.png",
        "Sequência UC03 - Solicitar matrícula",
        ["Aluno", "Interface Web", "MatriculaService", "PlanoRepository", "Banco de dados"],
        [
            (0, 1, "Seleciona plano"),
            (1, 2, "envia solicitação"),
            (2, 3, "consulta disponibilidade"),
            (3, 4, "busca plano ativo"),
            (4, 3, "retorna dados"),
            (2, 4, "grava matrícula e protocolo"),
            (2, 1, "confirma solicitação"),
            (1, 0, "exibe protocolo"),
        ],
    )
    booking = sequence_diagram(
        "figura-08-sequencia-agendamento.png",
        "Sequência UC04 - Agendar aula",
        ["Aluno", "Interface Web", "AgendaService", "ReservaRepository", "Banco de dados"],
        [
            (0, 1, "Solicita reserva"),
            (1, 2, "envia aula escolhida"),
            (2, 3, "verifica conflito e duplicidade"),
            (3, 4, "consulta reservas atuais"),
            (4, 3, "retorna ocupação"),
            (2, 4, "salva reserva"),
            (2, 1, "retorna vagas atualizadas"),
            (1, 0, "mostra confirmação"),
        ],
    )
    workout = sequence_diagram(
        "figura-09-sequencia-treino.png",
        "Sequência UC06 - Registrar treino",
        ["Aluno", "Interface Web", "TreinoService", "EvolucaoRepository", "Banco de dados"],
        [
            (0, 1, "Marca exercícios concluídos"),
            (1, 2, "envia registro"),
            (2, 2, "valida cargas e repetições"),
            (2, 3, "cria registro de treino"),
            (3, 4, "persiste histórico"),
            (4, 3, "confirma gravação"),
            (2, 1, "calcula evolução"),
            (1, 0, "atualiza dashboard"),
        ],
    )
    return [enrollment, booking, workout]


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color.replace("#", ""))
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]]) -> None:
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        set_cell_text(header_cells[index], header, True)
        set_cell_fill(header_cells[index], "2F3742")
        header_cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in header_cells[index].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_border(header_cells[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell_index, text in enumerate(row):
            set_cell_text(cells[cell_index], text)
            cells[cell_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_fill(cells[cell_index], "F4F6F2")
            set_cell_border(cells[cell_index])
    source = doc.add_paragraph("Fonte: elaboracao propria (2026).")
    source.runs[0].font.size = Pt(8)
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_figure(doc: Document, image_path: Path, caption: str, width_inches: float = 6.6) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Academia Conectada - Entrega 2 - pagina ")
    run.font.size = Pt(8)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_sep)
    run._r.append(text)
    run._r.append(fld_char_end)


def remove_paragraph_borders(style) -> None:
    p_pr = style.element.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag == qn("w:pBdr"):
            p_pr.remove(child)


ACCENT_REPLACEMENTS = [
    ("Analise", "Análise"),
    ("analise", "análise"),
    ("Informacoes", "Informações"),
    ("informacoes", "informações"),
    ("Introducao", "Introdução"),
    ("introducao", "introdução"),
    ("versao", "versão"),
    ("Versao", "Versão"),
    ("documentacao", "documentação"),
    ("Documentacao", "Documentação"),
    ("aplicacao", "aplicação"),
    ("Aplicacao", "Aplicação"),
    ("repositorio", "repositório"),
    ("Repositorio", "Repositório"),
    ("codigo", "código"),
    ("codigo fonte", "código fonte"),
    ("acompanhamento", "acompanhamento"),
    ("Interessados", "Interessados"),
    ("publica", "pública"),
    ("Publica", "Pública"),
    ("publicas", "públicas"),
    ("Conteudo", "Conteúdo"),
    ("conteudo", "conteúdo"),
    ("Matricula", "Matrícula"),
    ("matricula", "matrícula"),
    ("situacao", "situação"),
    ("participacao", "participação"),
    ("evolucao", "evolução"),
    ("Evolucao", "Evolução"),
    ("execucao", "execução"),
    ("gestao", "gestão"),
    ("Gestao", "Gestão"),
    ("Administracao", "Administração"),
    ("administracao", "administração"),
    ("funcionais", "funcionais"),
    ("nao", "não"),
    ("Nao", "Não"),
    ("Descricao", "Descrição"),
    ("descricao", "descrição"),
    ("sequencia", "sequência"),
    ("Sequencia", "Sequência"),
    ("dominio", "domínio"),
    ("Dominio", "Domínio"),
    ("Prototipo", "Protótipo"),
    ("Prototipos", "Protótipos"),
    ("prototipo", "protótipo"),
    ("prototipos", "protótipos"),
    ("academica", "acadêmica"),
    ("academico", "acadêmico"),
    ("Academico", "Acadêmico"),
    ("bancarios", "bancários"),
    ("Area", "Área"),
    ("areas", "áreas"),
    ("area", "área"),
    ("compativel", "compatível"),
    ("incompativel", "incompatível"),
    ("disponivel", "disponível"),
    ("indisponivel", "indisponível"),
    ("temporariamente", "temporariamente"),
    ("horarios", "horários"),
    ("Horario", "Horário"),
    ("horario", "horário"),
    ("localizacao", "localização"),
    ("saida", "saída"),
    ("recuperacao", "recuperação"),
    ("criterio", "critério"),
    ("Criterio", "Critério"),
    ("aceitacao", "aceitação"),
    ("previo", "prévio"),
    ("utilizavel", "utilizável"),
    ("rotulos", "rótulos"),
    ("ate", "até"),
    ("Paginas", "Páginas"),
    ("paginas", "páginas"),
    ("movel", "móvel"),
    ("estavel", "estável"),
    ("conexao", "conexão"),
    ("Seguranca", "Segurança"),
    ("seguranca", "segurança"),
    ("Privacidade", "Privacidade"),
    ("necessarios", "necessários"),
    ("correcao", "correção"),
    ("exclusao", "exclusão"),
    ("versoes", "versões"),
    ("homologacao", "homologação"),
    ("Confiabilidade", "Confiabilidade"),
    ("duplicidade", "duplicidade"),
    ("associacoes", "associações"),
    ("Associacao", "Associação"),
    ("consistencia", "consistência"),
    ("Consideracoes", "Considerações"),
    ("consideracoes", "considerações"),
    ("Alem", "Além"),
    ("alem", "além"),
    ("persistencia", "persistência"),
    ("sessao", "sessão"),
    ("Sessoes", "Sessões"),
    ("sessoes", "sessões"),
    ("series", "séries"),
    ("exercicios", "exercícios"),
    ("Exercicio", "Exercício"),
    ("orientacoes", "orientações"),
    ("observacoes", "observações"),
    ("validacao", "validação"),
    ("autorizacao", "autorização"),
    ("autenticacao", "autenticação"),
    ("Autenticacao", "Autenticação"),
    ("responsavel", "responsável"),
    ("Beneficios", "Benefícios"),
    ("beneficios", "benefícios"),
    ("historico", "histórico"),
    ("Historico", "Histórico"),
    ("historica", "histórica"),
    ("propria", "própria"),
    ("elaboracao", "elaboração"),
    ("Identificacao", "Identificação"),
    ("identificacao", "identificação"),
    ("Instituicao", "Instituição"),
    ("instituicao", "instituição"),
    ("Informacao", "Informação"),
    ("informacao", "informação"),
    ("informacoes", "informações"),
    ("Numero", "Número"),
    ("Usuario", "Usuário"),
    ("usuario", "usuário"),
    ("Descricao verificavel", "Descrição verificável"),
    ("verificavel", "verificável"),
    ("concluidos", "concluídos"),
    ("solicitacao", "solicitação"),
    ("excecoes", "exceções"),
    ("ja", "já"),
    ("invalidas", "inválidas"),
    ("invalido", "inválido"),
    ("modulo", "módulo"),
    ("obrigatorio", "obrigatório"),
    ("obrigatoria", "obrigatória"),
    ("vinculo", "vínculo"),
    ("acao", "ação"),
    ("Acao", "Ação"),
    ("acoes", "ações"),
    ("Acoes", "Ações"),
    ("Secao", "Seção"),
    ("secao", "seção"),
    ("Botoes", "Botões"),
    ("botoes", "botões"),
    ("botao", "botão"),
    ("Botao", "Botão"),
    ("metricas", "métricas"),
    ("manutencao", "manutenção"),
    ("organizacao", "organização"),
    ("separacao", "separação"),
    ("interacoes", "interações"),
    ("servicos", "serviços"),
    ("colaboracao", "colaboração"),
    ("repositorios", "repositórios"),
    ("proximas", "próximas"),
    ("tecnicas", "técnicas"),
    ("sao", "são"),
    ("revisao", "revisão"),
    ("interrupcao", "interrupção"),
    ("disponibilidade", "disponibilidade"),
    ("confirmacao", "confirmação"),
    ("gravacao", "gravação"),
    ("repeticoes", "repetições"),
    ("prescricao", "prescrição"),
]


def accent_text(value: str) -> str:
    text = value
    for old, new in ACCENT_REPLACEMENTS:
        text = re.sub(rf"\b{re.escape(old)}\b", new, text)
    return text


def accent_document(doc: Document) -> None:
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            has_break = bool(run._r.xpath(".//w:br"))
            has_field = bool(run._r.xpath(".//w:fldChar")) or bool(run._r.xpath(".//w:instrText"))
            if run.text and not has_break and not has_field:
                run.text = accent_text(run.text)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(32, 36, 44)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size in [("Title", 19), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 11)]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(7)
        remove_paragraph_borders(style)


def add_heading(doc: Document, number: str, title: str, level: int = 1) -> None:
    if number:
        doc.add_heading(f"{number} {title}", level=level)
    else:
        doc.add_heading(title, level=level)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def create_docx(images: dict[str, Path]) -> None:
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Academia Conectada").bold = True
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Sistema Web para Gestao de Alunos, Planos, Treinos e Aulas")
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.bold = True
    doc.add_paragraph()
    meta = [
        ("Componente curricular", "Prat. Prof. em Analise Desenvolvimento Sistemas"),
        ("Entrega", "Segunda versao da documentacao e interface implementada"),
        ("Instituicao", "Universidade Presbiteriana Mackenzie"),
        ("Data", "05 de setembro de 2026"),
        ("Aplicacao publicada", "https://academia-conectada.vercel.app"),
        ("Repositorio", "https://github.com/Lanpaiva/academia-conectada"),
        ("Quadro", "https://github.com/users/Lanpaiva/projects/2"),
    ]
    add_table(doc, "Tabela 1 - Identificacao do projeto", ["Campo", "Informacao"], meta)
    doc.add_paragraph(
        "Alan Araujo Paiva, Ian Avila Guerra, Julio de Moura Stelzer e Matheus Araujo dos Santos Nunes."
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    add_heading(doc, "", "Sumario")
    for item in [
        "1 Introducao",
        "2 Informacoes do projeto",
        "3 Objetivos funcionais",
        "4 Objetivos nao funcionais",
        "5 Diagrama de casos de uso",
        "6 Descricao detalhada dos casos de uso principais",
        "7 Prototipos de tela",
        "8 Interface implementada",
        "9 Modelo de dominio",
        "10 Diagrama de classes de projeto",
        "11 Diagramas de sequencia de projeto",
        "12 Consideracoes finais",
    ]:
        doc.add_paragraph(item)
    doc.add_page_break()

    add_heading(doc, "", "Lista de figuras")
    figure_list = [
        "Figura 1 - Diagrama geral de casos de uso do sistema Academia Conectada",
        "Figura 2 - Prototipo da tela publica de planos e aulas",
        "Figura 3 - Prototipo do dashboard do aluno",
        "Figura 4 - Prototipo das areas de instrutor e administracao",
        "Figura 5 - Modelo de dominio do sistema",
        "Figura 6 - Diagrama de classes de projeto",
        "Figura 7 - Diagrama de sequencia do fluxo de matricula",
        "Figura 8 - Diagrama de sequencia do fluxo de agendamento",
        "Figura 9 - Diagrama de sequencia do fluxo de registro de treino",
    ]
    for item in figure_list:
        doc.add_paragraph(item)

    add_heading(doc, "", "Lista de tabelas")
    table_list = [
        "Tabela 1 - Identificacao do projeto",
        "Tabela 2 - Integrantes do grupo",
        "Tabela 3 - Interessados no projeto",
        "Tabela 4 - Objetivos funcionais do sistema",
        "Tabela 5 - Objetivos nao funcionais do sistema",
        "Tabela 6 - Rastreabilidade entre casos de uso e interface",
        "Tabela 7 - Entidades do modelo de dominio",
    ]
    for item in table_list:
        doc.add_paragraph(item)
    doc.add_page_break()

    add_heading(doc, "1", "Introducao")
    doc.add_paragraph(
        "O projeto Academia Conectada propõe um sistema web para apoiar a rotina de uma academia, "
        "centralizando informacoes de alunos, planos, aulas, reservas, treinos e evolucao. Esta segunda "
        "versao revisa os elementos da entrega inicial e acrescenta prototipos de tela, modelo de dominio, "
        "diagrama de classes de projeto e diagramas de sequencia para os principais fluxos."
    )
    doc.add_paragraph(
        "Alem da documentacao, foi desenvolvida uma primeira interface executavel em Next.js. A interface "
        "usa dados demonstrativos e estado local para representar os fluxos principais enquanto o projeto "
        "ainda nao possui banco de dados, autenticacao real ou persistencia de informacoes."
    )

    add_heading(doc, "2", "Informacoes do projeto")
    add_heading(doc, "2.1", "Titulo do projeto", 2)
    doc.add_paragraph("Academia Conectada - Sistema Web para Gestao de Alunos, Planos, Treinos e Aulas.")
    add_heading(doc, "2.2", "Integrantes do grupo", 2)
    add_table(
        doc,
        "Tabela 2 - Integrantes do grupo",
        ["Numero", "Nome", "RA"],
        [
            ["1", "Alan Araujo Paiva", "10423944"],
            ["2", "Ian Avila Guerra", "Pendente"],
            ["3", "Julio de Moura Stelzer", "10336198"],
            ["4", "Matheus Araujo dos Santos Nunes", "10444279"],
        ],
    )
    add_heading(doc, "2.3", "URL do repositorio de codigo fonte", 2)
    doc.add_paragraph("https://github.com/Lanpaiva/academia-conectada")
    add_heading(doc, "2.4", "URL do quadro de acompanhamento", 2)
    doc.add_paragraph("https://github.com/users/Lanpaiva/projects/2")
    add_heading(doc, "2.5", "Interessados", 2)
    add_table(
        doc,
        "Tabela 3 - Interessados no projeto",
        ["Interessado", "Interesse principal"],
        [
            ["Visitantes", "Conhecer planos, modalidades, horarios, localizacao e canais de contato."],
            ["Alunos", "Solicitar matricula, reservar aulas, consultar treino e registrar evolucao."],
            ["Instrutores", "Criar, revisar e acompanhar planos de treino dos alunos."],
            ["Administracao da academia", "Manter cadastros, planos, modalidades, horarios e aulas."],
        ],
    )

    add_heading(doc, "3", "Objetivos funcionais")
    add_table(
        doc,
        "Tabela 4 - Objetivos funcionais do sistema",
        ["ID", "Objetivo", "Descricao verificavel"],
        [
            ["OF01", "Autenticacao", "Permitir cadastro, entrada, saida e recuperacao de acesso conforme perfil."],
            ["OF02", "Conteudo publico", "Exibir academia, modalidades, planos, horarios, localizacao e contato."],
            ["OF03", "Matricula", "Permitir que o aluno solicite matricula em um plano e acompanhe sua situacao."],
            ["OF04", "Agenda de aulas", "Permitir consultar disponibilidade, reservar vaga e cancelar participacao."],
            ["OF05", "Plano de treino", "Exibir exercicios, series, repeticoes, carga, descanso e orientacoes."],
            ["OF06", "Evolucao", "Registrar execucao do treino, cargas utilizadas e medidas de acompanhamento."],
            ["OF07", "Gestao de treinos", "Permitir ao instrutor criar, atualizar e atribuir planos de treino."],
            ["OF08", "Administracao", "Permitir manter alunos, planos, modalidades, horarios e aulas."],
        ],
    )

    doc.add_page_break()
    add_heading(doc, "4", "Objetivos nao funcionais")
    add_table(
        doc,
        "Tabela 5 - Objetivos nao funcionais do sistema",
        ["ID", "Atributo", "Criterio de aceitacao"],
        [
            ["ONF01", "Usabilidade", "Fluxos principais concluidos sem treinamento previo e com mensagens claras."],
            ["ONF02", "Responsividade", "Interface utilizavel entre 360 px e 1920 px sem rolagem horizontal indevida."],
            ["ONF03", "Acessibilidade", "Buscar conformidade com WCAG 2.1 AA: contraste, teclado, rotulos e textos alternativos."],
            ["ONF04", "Desempenho", "Paginas publicas carregando em ate 3 segundos em conexao movel estavel nos testes."],
            ["ONF05", "Seguranca", "Usar HTTPS, validacao de entrada, controle por perfil e senhas sem texto puro."],
            ["ONF06", "Privacidade", "Coletar apenas dados necessarios e permitir correcao ou exclusao conforme a LGPD."],
            ["ONF07", "Compatibilidade", "Suportar versoes atuais de Chrome, Edge, Firefox e Safari durante homologacao."],
            ["ONF08", "Confiabilidade", "Impedir duplicidade de agendamentos e preservar consistencia de vagas."],
        ],
    )

    doc.add_page_break()
    add_heading(doc, "5", "Diagrama de casos de uso")
    doc.add_paragraph(
        "O diagrama de casos de uso apresenta os atores externos e as principais capacidades do sistema. "
        "As associacoes indicam quais perfis participam de cada caso, sem representar a ordem temporal dos passos."
    )
    add_figure(doc, images["use_cases"], "Figura 1 - Diagrama geral de casos de uso do sistema Academia Conectada")

    add_heading(doc, "6", "Descricao detalhada dos casos de uso principais")
    use_case_rows = [
        (
            "UC01 - Cadastrar-se e entrar",
            "Visitante, aluno, instrutor e administrador.",
            "Criar conta ou autenticar usuario existente e direcionar para a area compatível com seu perfil.",
            "Usuario informa dados, sistema valida obrigatoriedade, unicidade e credenciais, cria sessao e redireciona.",
            "E-mail ja cadastrado, credenciais invalidas ou conta desabilitada.",
        ),
        (
            "UC02 - Consultar planos e modalidades",
            "Visitante e aluno.",
            "Apresentar informacoes claras para apoiar a escolha de servicos da academia.",
            "Sistema lista planos, modalidades e horarios; usuario filtra, consulta detalhes e inicia contato ou matricula.",
            "Nenhum item disponivel ou informacao temporariamente indisponivel.",
        ),
        (
            "UC03 - Solicitar matricula",
            "Aluno.",
            "Registrar escolha de plano sem armazenar dados bancarios reais nesta etapa academica.",
            "Aluno escolhe plano, confirma dados, aceita termos, sistema grava solicitacao e exibe protocolo.",
            "Dados incompletos, plano indisponivel ou matricula ativa incompatível.",
        ),
        (
            "UC04 - Agendar ou cancelar aula",
            "Aluno.",
            "Reservar ou liberar vaga em aula coletiva com consistencia de capacidade.",
            "Aluno seleciona aula, sistema valida matricula, conflito, duplicidade e disponibilidade, depois confirma reserva.",
            "Aula lotada, conflito de horario, duplicidade ou cancelamento fora do prazo.",
        ),
        (
            "UC05 - Consultar plano de treino",
            "Aluno e instrutor.",
            "Disponibilizar a prescricao vigente de forma clara e organizada.",
            "Sistema identifica plano vigente, usuario escolhe sessao e consulta exercicios, series, cargas e descanso.",
            "Plano inexistente, vencido ou acesso sem autorizacao.",
        ),
        (
            "UC06 - Registrar treino e evolucao",
            "Aluno.",
            "Registrar execucao do treino e dados de acompanhamento para consulta historica.",
            "Aluno marca exercicios, informa cargas e observacoes, sistema valida e salva o historico.",
            "Valor invalido ou interrupcao do preenchimento.",
        ),
        (
            "UC07 - Gerenciar treinos de alunos",
            "Instrutor.",
            "Criar, revisar, atribuir e encerrar planos de treino dos alunos autorizados.",
            "Instrutor seleciona aluno, define sessoes e exercicios, valida campos e publica nova versao.",
            "Plano incompleto, aluno nao autorizado ou revisao sem apagar historico.",
        ),
        (
            "UC08 - Administrar cadastros, planos e aulas",
            "Administrador.",
            "Manter dados de alunos, modalidades, planos, horarios e aulas utilizados pelo sistema.",
            "Administrador escolhe modulo, cria ou altera registro, sistema valida e apresenta lista atualizada.",
            "Dado obrigatorio ausente, registro duplicado ou exclusao impedida por vinculo existente.",
        ),
    ]
    for uc, actors, objective, flow, alternatives in use_case_rows:
        add_heading(doc, "", uc, 2)
        doc.add_paragraph(f"Atores: {actors}")
        doc.add_paragraph(f"Objetivo: {objective}")
        doc.add_paragraph(f"Fluxo principal: {flow}")
        doc.add_paragraph(f"Fluxos alternativos e excecoes: {alternatives}")

    add_heading(doc, "7", "Prototipos de tela")
    doc.add_paragraph(
        "Os prototipos representam a organizacao visual prevista para os perfis do sistema. Eles foram usados como "
        "referencia direta para a interface executavel publicada."
    )
    add_figure(doc, images["prototype_public"], "Figura 2 - Prototipo da tela publica de planos e aulas")
    add_figure(doc, images["prototype_student"], "Figura 3 - Prototipo do dashboard do aluno")
    add_figure(doc, images["prototype_management"], "Figura 4 - Prototipo das areas de instrutor e administracao")

    doc.add_page_break()
    add_heading(doc, "8", "Interface implementada")
    doc.add_paragraph(
        "A interface inicial foi implementada com Next.js, React, TypeScript, Tailwind CSS e Lucide React. "
        "A versao publicada permite alternar entre perfis e executar interacoes demonstrativas de matricula, "
        "reserva e cancelamento de aula, registro de treino, consulta de evolucao, gestao de alunos pelo instrutor "
        "e manutencao administrativa."
    )
    doc.add_paragraph("URL da aplicacao: https://academia-conectada.vercel.app")
    add_table(
        doc,
        "Tabela 6 - Rastreabilidade entre casos de uso e interface",
        ["Caso de uso", "Elemento implementado na interface"],
        [
            ["UC01", "Painel de acesso com e-mail, perfil e acao Entrar."],
            ["UC02", "Secao de planos e agenda publica para visitante."],
            ["UC03", "Botoes de solicitacao de matricula com protocolo demonstrativo."],
            ["UC04", "Lista de aulas com controle visual de vagas e botao reservar ou cancelar."],
            ["UC05", "Secao Meu treino com exercicios, series, cargas e descanso."],
            ["UC06", "Checkboxes de execucao, progresso percentual e botao Salvar registro."],
            ["UC07", "Area do instrutor com alunos, status e acao Editar treino."],
            ["UC08", "Painel administrativo com metricas e manutencao de cadastros."],
        ],
    )

    doc.add_page_break()
    add_heading(doc, "9", "Modelo de dominio")
    doc.add_paragraph(
        "O modelo de dominio organiza os conceitos principais da academia e seus relacionamentos. A separacao entre "
        "usuario, perfis, matricula, reserva, aula e treino facilita a evolucao para persistencia com PostgreSQL e Prisma."
    )
    add_figure(doc, images["domain"], "Figura 5 - Modelo de dominio do sistema")
    add_table(
        doc,
        "Tabela 7 - Entidades do modelo de dominio",
        ["Entidade", "Responsabilidade"],
        [
            ["Usuario", "Representar credenciais, identificacao e perfil de acesso."],
            ["Aluno", "Representar o usuario que contrata plano, agenda aula e registra evolucao."],
            ["Instrutor", "Prescrever e revisar planos de treino."],
            ["Administrador", "Manter cadastros e regras operacionais."],
            ["Matricula", "Vincular aluno a plano contratado e controlar situacao."],
            ["Plano", "Definir beneficios, valor demonstrativo e regras de acesso."],
            ["Aula", "Representar horario, modalidade, instrutor, local e capacidade."],
            ["Reserva", "Vincular aluno a aula e controlar confirmacao ou cancelamento."],
            ["Plano de treino", "Organizar sessoes e exercicios prescritos pelo instrutor."],
            ["Registro de evolucao", "Guardar historico de treino, medidas e observacoes."],
        ],
    )

    add_heading(doc, "10", "Diagrama de classes de projeto")
    doc.add_paragraph(
        "O diagrama de classes de projeto traduz o modelo de dominio para estruturas que poderao ser implementadas "
        "em TypeScript e posteriormente persistidas com Prisma."
    )
    add_figure(doc, images["classes"], "Figura 6 - Diagrama de classes de projeto", width_inches=6.7)

    doc.add_page_break()
    add_heading(doc, "11", "Diagramas de sequencia de projeto")
    doc.add_paragraph(
        "Os diagramas de sequencia detalham a colaboracao entre usuario, interface, servicos de aplicacao, repositorios "
        "e banco de dados nos fluxos principais."
    )
    add_figure(doc, images["sequence_enrollment"], "Figura 7 - Diagrama de sequencia do fluxo de matricula")
    doc.add_page_break()
    add_figure(doc, images["sequence_booking"], "Figura 8 - Diagrama de sequencia do fluxo de agendamento")
    doc.add_page_break()
    add_figure(doc, images["sequence_workout"], "Figura 9 - Diagrama de sequencia do fluxo de registro de treino")

    add_heading(doc, "12", "Consideracoes finais")
    doc.add_paragraph(
        "A segunda versao consolida a documentacao revisada, apresenta os novos artefatos de modelagem e entrega uma "
        "interface inicial publicada. As proximas etapas tecnicas recomendadas sao implementar persistencia, autenticacao "
        "real, regras de autorizacao por perfil e testes automatizados para os fluxos mais importantes."
    )

    accent_document(doc)
    doc.save(DOCX_PATH)


def main() -> None:
    ensure_dirs()
    sequence = sequence_images()
    images = {
        "use_cases": use_case_diagram(),
        "prototype_public": prototype_public(),
        "prototype_student": prototype_student(),
        "prototype_management": prototype_management(),
        "domain": domain_model(),
        "classes": class_diagram(),
        "sequence_enrollment": sequence[0],
        "sequence_booking": sequence[1],
        "sequence_workout": sequence[2],
    }
    create_docx(images)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
